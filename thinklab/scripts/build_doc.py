#!/usr/bin/env python3
"""ThinkLab document builder.

Turns a stage JSON file into an Arabic RTL Word document styled like the
reference documents (pink/cyan/red/yellow vocabulary), then converts it to PDF
with LibreOffice.

Usage:
    python build_doc.py stage.json [more.json ...] [--out DIR] [--no-pdf]

See ../references/json-format.md for the block types.
"""
import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    subprocess.run([sys.executable, "-m", "pip", "install", "-q", "python-docx"], check=False)
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

SKILL_DIR = Path(__file__).resolve().parent.parent
FONTS_DIR = SKILL_DIR / "assets" / "fonts"

# Colours taken from the reference documents.
PINK = "FF0066"
CYAN = "00B0F0"
RED = "FF0000"
DARK_RED = "C00000"
YELLOW = "FFFF00"
WHITE = "FFFFFF"
BLACK = "000000"
HEADING_BLUE = "365F91"
GREEN = "00B050"
# "role" theme colours, from the user's per-role reference ("حارس البوابة — المميزات ورحلة المستخدم").
ROLE_BLUE = "2E74B5"
ROLE_DARK_BLUE = "1F4D78"
ROLE_GREY = "EFEFEF"
ROLE_RULE = "BFBFBF"

AR_FONT = "Amiri"
LATIN_FONT = "Times New Roman"
BODY_SIZE = 12

# theme -> (Arabic font, Latin font, body size, page margin in cm, default paragraph alignment)
THEMES = {
    "classic": ("Amiri", "Times New Roman", 12, 1.27, "both"),
    "role": ("Noto Sans Arabic", "Arial", 10.5, 1.76, "right"),
}


def apply_theme(name):
    global AR_FONT, LATIN_FONT, BODY_SIZE
    AR_FONT, LATIN_FONT, BODY_SIZE, margin, align = THEMES[name]
    return margin, align

ARABIC_RE = re.compile(r"[؀-ۿݐ-ݿﭐ-﷿ﹰ-﻿]")
LATIN_CHUNK_RE = re.compile(r"([A-Za-z][A-Za-z0-9&/+\-_'.]*(?:[ ]+[A-Za-z0-9&/+\-_'.]+)*)")
INLINE_RE = re.compile(r"\[\[(\w+):(.+?)\]\]|\*\*(.+?)\*\*")

ORDINALS = ["الأولى", "الثانية", "الثالثة", "الرابعة", "الخامسة", "السادسة", "السابعة",
            "الثامنة", "التاسعة", "العاشرة", "الحادية عشرة", "الثانية عشرة"]


# ---------------------------------------------------------------- low level XML

def _set(el, tag, **attrs):
    child = el.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        el.append(child)
    for k, v in attrs.items():
        child.set(qn(k), str(v))
    return child


def para_rtl(p, align="right"):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    pPr.insert(0, bidi)
    # In a bidi paragraph "left" is the start side (visually right) and vice versa.
    jc = {"right": "left", "left": "right"}.get(align, align)
    _set(pPr, "w:jc", **{"w:val": jc})
    return p


def para_spacing(p, before=0, after=2, line=1.1):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def para_shading(p, fill):
    pPr = p._p.get_or_add_pPr()
    _set(pPr, "w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill})


def para_indent_start(p, start_twips, hanging_twips=0):
    pPr = p._p.get_or_add_pPr()
    attrs = {"w:right": start_twips}
    if hanging_twips:
        attrs["w:hanging"] = hanging_twips
    _set(pPr, "w:ind", **attrs)


def add_run(p, text, color=None, bold=False, size=None, bg=None, highlight=None,
            font=None, rtl=None, italic=False):
    run = p.add_run(text)
    rPr = run._r.get_or_add_rPr()
    f = font or AR_FONT
    latin = font or LATIN_FONT
    rf = _set(rPr, "w:rFonts")
    rf.set(qn("w:ascii"), latin)
    rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:cs"), f)
    rf.set(qn("w:eastAsia"), latin)
    sz = int((size or BODY_SIZE) * 2)
    if bold:
        _set(rPr, "w:b")
        _set(rPr, "w:bCs")
    if italic:
        _set(rPr, "w:i")
        _set(rPr, "w:iCs")
    if color:
        _set(rPr, "w:color", **{"w:val": color})
    _set(rPr, "w:sz", **{"w:val": sz})
    _set(rPr, "w:szCs", **{"w:val": sz})
    if highlight:
        _set(rPr, "w:highlight", **{"w:val": highlight})
    if bg:
        _set(rPr, "w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": bg})
    if rtl is None:
        rtl = bool(ARABIC_RE.search(text)) or not re.search(r"[A-Za-z]", text)
    if rtl:
        _set(rPr, "w:rtl")
    return run


def add_text(p, text, **style):
    """Add text, splitting Latin chunks into LTR runs so mixed lines read correctly."""
    for i, part in enumerate(LATIN_CHUNK_RE.split(text)):
        if not part:
            continue
        add_run(p, part, rtl=(i % 2 == 0), **style)


INLINE_STYLES = {
    "red": {"color": RED},
    "darkred": {"color": DARK_RED},
    "cyan": {"color": CYAN},
    "blue": {"color": CYAN},
    "pink": {"color": PINK},
    "green": {"color": GREEN},
    "pinkbg": {"color": WHITE, "bg": PINK},
    "yellowbg": {"bg": YELLOW},
    "hl": {"highlight": "cyan"},
    "mark": {"highlight": "yellow"},
    "b": {"bold": True},
}


def add_rich(p, text, **base):
    """Inline markup: **bold**, [[red:..]], [[cyan:..]], [[pink:..]], [[pinkbg:..]], [[hl:..]] ..."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            add_text(p, text[pos:m.start()], **base)
        if m.group(3) is not None:
            add_text(p, m.group(3), **{**base, "bold": True})
        else:
            style = INLINE_STYLES.get(m.group(1).lower(), {})
            add_text(p, m.group(2), **{**base, **style})
        pos = m.end()
    if pos < len(text):
        add_text(p, text[pos:], **base)


# ---------------------------------------------------------------- blocks

class Builder:
    def __init__(self, data):
        self.data = data
        self.doc = Document()
        self.error_no = 0
        self.success_no = 0
        self.card_no = 0
        self.margin, self.default_align = apply_theme(data.get("theme", "classic"))
        self._setup()

    def _setup(self):
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(sec, side, Cm(self.margin))
        _set(sec._sectPr, "w:bidi")
        st = self.doc.styles["Normal"]
        st.font.name = LATIN_FONT
        st.font.size = Pt(BODY_SIZE)
        rf = _set(st.element.get_or_add_rPr(), "w:rFonts")
        rf.set(qn("w:cs"), AR_FONT)
        _set(st.element.get_or_add_rPr(), "w:szCs", **{"w:val": BODY_SIZE * 2})
        lang = _set(st.element.get_or_add_rPr(), "w:lang")
        lang.set(qn("w:bidi"), "ar-EG")
        props = self.doc.core_properties
        props.title = self.data.get("title", "")
        props.author = "ThinkLab"

    def p(self, align="both", after=2, before=0, line=1.1):
        p = self.doc.add_paragraph()
        para_rtl(p, align)
        para_spacing(p, before, after, line)
        return p

    # --- shared ---
    def b_basmala(self, b):
        p = self.p("center", after=4)
        add_text(p, b.get("text", "بسم الله الرحمن الرحيم"), size=13)

    def b_banner(self, b):
        style = b.get("style", "yellow")
        conf = {
            "yellow": (YELLOW, RED, 13, "center"),
            "cyan": (CYAN, PINK, 18, "center"),
            "cyan_white": (CYAN, WHITE, 13, "center"),
            "pink": (PINK, WHITE, 14, "center"),
            "title": (YELLOW, BLACK, 13, "right"),
        }[style]
        fill, color, size, align = conf
        p = self.p(align, before=b.get("before", 6), after=4, line=1.0)
        para_shading(p, fill)
        p.paragraph_format.keep_with_next = True
        add_rich(p, b["text"], color=color, size=size, bold=b.get("bold", False))

    def b_paragraph(self, b):
        role = self.default_align == "right"
        p = self.p(b.get("align", self.default_align), after=b.get("after", 3 if role else 2),
                   line=1.0 if role else 1.1)
        add_rich(p, b["text"], size=b.get("size"), color=b.get("color"))

    def b_spacer(self, b):
        for _ in range(int(b.get("lines", 1))):
            self.p(after=0)

    # --- template 00: considerations ---
    def b_point(self, b):
        p = self.p("both", after=3)
        add_run(p, "ؐ ", color=CYAN, size=14)
        add_rich(p, b["text"])
        for line in b.get("lines", []):
            q = self.p("both", after=2)
            add_rich(q, line)

    def b_item(self, b):
        p = self.p("both", after=4)
        add_run(p, "֎ ", color=PINK, size=14, font="DejaVu Sans")
        add_rich(p, b["label"] + " : ", color=PINK)
        add_rich(p, b.get("text", ""))
        for line in b.get("lines", []):
            q = self.p("both", after=2)
            add_rich(q, line)

    # --- template 01: structure ---
    def b_intro(self, b):
        p = self.p("both", after=3)
        add_run(p, "❒ ", color=RED, font="DejaVu Sans")
        add_rich(p, b["label"] + " : ", color=RED)
        add_rich(p, b.get("text", ""))
        for line in b.get("lines", []):
            q = self.p("both", after=2)
            add_rich(q, line)

    def b_module(self, b):
        p = self.p("both", after=1, before=2)
        add_run(p, "֎ ", color=PINK, font="DejaVu Sans")
        add_rich(p, b["label"] + " :", color=WHITE, bg=PINK)
        if b.get("text"):
            add_rich(p, " " + b["text"])
        if b.get("list"):
            q = self.p("both", after=1)
            add_rich(q, " ".join("- " + x.rstrip(".") + "." for x in b["list"]))
        for line in b.get("lines", []):
            q = self.p("both", after=1)
            add_rich(q, line)
        if b.get("goal"):
            self.b_goal({"text": b["goal"]})

    def b_goal(self, b):
        p = self.p("both", after=3)
        add_rich(p, "### الهدف منه: " + b["text"])

    def b_must(self, b):
        p = self.p("both", after=3)
        add_rich(p, b.get("label", "يجب مراعاة") + " ", color=DARK_RED)
        add_rich(p, b["text"])

    def b_journey(self, b):
        p = self.p("both", after=2)
        add_run(p, "❒ ", color=RED, font="DejaVu Sans")
        add_text(p, "رحلة ")
        add_rich(p, b["name"], color=CYAN)
        add_rich(p, " " + b.get("text", ""))
        for line in b.get("lines", []):
            q = self.p("both", after=1)
            add_rich(q, line)
        if b.get("buttons"):
            q = self.p("both", after=2)
            add_rich(q, "والأزرار: " + " ".join("- " + x.rstrip(".") + "." for x in b["buttons"]))

    # --- template 02/03: feature cards ---
    def b_card(self, b):
        self.card_no += 1
        title = b["title"]
        if not re.match(r"(?i)card\s*\d+", title):
            title = f"Card {self.card_no}: {title}"
        p = self.p("right", before=8, after=2, line=1.0)
        p.paragraph_format.keep_with_next = True
        add_text(p, title, size=13, highlight="cyan" if b.get("style") == "cyan" else None,
                 bg=None if b.get("style") == "cyan" else YELLOW, bold=False)
        for it in b.get("items", []):
            hl = False
            if isinstance(it, dict):
                text, hl = it["text"], it.get("highlight", False)
            else:
                text = it
                if text.startswith("!"):
                    text, hl = text[1:].strip(), True
            q = self.p("right", after=0, line=1.05)
            add_run(q, "• ", size=11)
            add_rich(q, text, size=11, highlight="cyan" if hl else None)

    # --- template 02 detailed journey ---
    def b_phase(self, b):
        text = b["text"]
        p = self.doc.add_paragraph()
        para_rtl(p, "right")
        para_spacing(p, before=14, after=8, line=1.0)
        add_rich(p, text, color=HEADING_BLUE, size=13)
        p.paragraph_format.keep_with_next = True

    def b_steps(self, b):
        for it in b["items"]:
            prefix, text, color = "", it, None
            if isinstance(it, dict):
                if "error" in it:
                    self.error_no += 1
                    prefix, text = f"حالة الخطأ ({self.error_no}): ", it["error"]
                    color = RED
                elif "success" in it:
                    prefix, text = "حالة النجاح: ", it["success"]
                    color = GREEN
                elif "note" in it:
                    prefix, text = "ملاحظة: ", it["note"]
                    color = DARK_RED
                else:
                    text = it.get("text", "")
            p = self.p("both", after=3, line=1.1)
            para_indent_start(p, 720, 360)
            add_run(p, "•\t")
            if prefix:
                add_rich(p, prefix, color=color)
            add_rich(p, text)

    # --- role file (reference: "حارس البوابة — المميزات ورحلة المستخدم") ---
    def _heading(self, text, color, size, before, after, italic=False, rule=False):
        p = self.p("right", before=before, after=after, line=1.0)
        p.paragraph_format.keep_with_next = True
        if rule:
            pBdr = _set(p._p.get_or_add_pPr(), "w:pBdr")
            _set(pBdr, "w:bottom", **{"w:val": "single", "w:sz": 6, "w:space": 4, "w:color": ROLE_RULE})
        add_rich(p, text, color=color, size=size, italic=italic)

    def b_h1(self, b):
        self._heading(b["text"], ROLE_BLUE, 16, b.get("before", 16), 8, rule=True)

    def b_h2(self, b):
        self._heading(b["text"], ROLE_BLUE, 13, 10, 4)

    def b_h3(self, b):
        self._heading(b["text"], ROLE_DARK_BLUE, 12, 8, 3)

    def b_h4(self, b):
        self._heading(b["text"], ROLE_BLUE, 11, 6, 2, italic=True)

    def b_meta(self, b):
        for row in b["items"]:
            p = self.p("right", after=1, line=1.0)
            add_rich(p, row["label"] + ": ", bold=True)
            add_rich(p, row["text"])

    def b_subtitle(self, b):
        p = self.p("right", before=4, after=2, line=1.0)
        p.paragraph_format.keep_with_next = True
        add_rich(p, b["text"], bold=True)

    def b_bullets(self, b):
        for it in b["items"]:
            p = self.p("right", after=0, line=1.0)
            para_indent_start(p, 540, 270)
            add_run(p, "•\t")
            add_rich(p, it)

    def b_divider(self, b):
        p = self.p("right", before=6, after=6, line=1.0)
        pBdr = _set(p._p.get_or_add_pPr(), "w:pBdr")
        _set(pBdr, "w:top", **{"w:val": "single", "w:sz": 4, "w:space": 1, "w:color": ROLE_RULE})

    # --- tables & sources ---
    def b_table(self, b):
        if b.get("title"):
            self.b_banner({"text": b["title"], "style": b.get("title_style", "cyan_white")})
        headers, rows = b["headers"], b["rows"]
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = t._tbl.tblPr
        _set(tblPr, "w:bidiVisual")
        hl_col = b.get("highlight_row")
        for r, row in enumerate([headers] + rows):
            trPr = t.rows[r]._tr.get_or_add_trPr()
            _set(trPr, "w:cantSplit")
            if r == 0:
                _set(trPr, "w:tblHeader")
            for c, val in enumerate(row):
                cell = t.cell(r, c)
                cell.text = ""
                p = cell.paragraphs[0]
                para_rtl(p, "center" if r == 0 else "right")
                para_spacing(p, 1, 1, 1.0)
                if r == 0:
                    grey = b.get("header_style") == "grey"
                    tcPr = cell._tc.get_or_add_tcPr()
                    _set(tcPr, "w:shd", **{"w:val": "clear", "w:color": "auto",
                                          "w:fill": ROLE_GREY if grey else CYAN})
                    add_rich(p, str(val), color=BLACK if grey else WHITE, bold=True,
                             size=BODY_SIZE if grey else 11)
                else:
                    if hl_col is not None and r - 1 == hl_col:
                        tcPr = cell._tc.get_or_add_tcPr()
                        _set(tcPr, "w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": "FFF2CC"})
                    add_rich(p, str(val), size=min(10, BODY_SIZE))
        self.p(after=2)

    def b_sources(self, b):
        self.b_banner({"text": b.get("title", "المصادر"), "style": "title"})
        for s in b["items"]:
            p = self.p("right", after=0, line=1.0)
            if isinstance(s, dict):
                add_rich(p, "• " + s.get("title", ""), size=10)
                if s.get("url"):
                    add_run(p, "  " + s["url"], size=9, color="0563C1", rtl=False)
            else:
                add_rich(p, "• " + s, size=10)

    def build(self):
        for b in self.data["blocks"]:
            fn = getattr(self, "b_" + b["type"], None)
            if fn is None:
                raise SystemExit(f"Unknown block type: {b['type']}")
            fn(b)
        return self.doc


# ---------------------------------------------------------------- conversion

def install_fonts():
    target = Path.home() / ".fonts"
    target.mkdir(parents=True, exist_ok=True)
    changed = False
    for f in FONTS_DIR.glob("*.ttf"):
        dst = target / f.name
        if not dst.exists():
            shutil.copy(f, dst)
            changed = True
    if changed and shutil.which("fc-cache"):
        subprocess.run(["fc-cache", "-f", str(target)], capture_output=True)


def find_soffice():
    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    for path in ("/usr/bin/soffice", "/usr/lib/libreoffice/program/soffice",
                 "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if os.path.exists(path):
            return path
    return None


def to_pdf(docx_path, out_dir):
    soffice = find_soffice()
    if not soffice:
        return None
    install_fonts()
    profile = Path(tempfile.gettempdir()) / "thinklab_lo_profile"
    cmd = [soffice, f"-env:UserInstallation=file://{profile}", "--headless",
           "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)]
    subprocess.run(cmd, capture_output=True, timeout=180)
    pdf = Path(out_dir) / (Path(docx_path).stem + ".pdf")
    return pdf if pdf.exists() else None


def safe_name(name):
    return re.sub(r'[\\/:*?"<>|]+', "-", name).strip() or "document"


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("inputs", nargs="+")
    ap.add_argument("--out", default=None, help="output directory (default: next to the JSON)")
    ap.add_argument("--no-pdf", action="store_true")
    args = ap.parse_args()

    for inp in args.inputs:
        data = json.loads(Path(inp).read_text(encoding="utf-8"))
        out_dir = Path(args.out or Path(inp).parent)
        out_dir.mkdir(parents=True, exist_ok=True)
        name = safe_name(data.get("file_name") or data.get("title") or Path(inp).stem)
        docx_path = out_dir / f"{name}.docx"
        Builder(data).build().save(docx_path)
        print(f"DOCX: {docx_path}")
        if args.no_pdf:
            continue
        pdf = to_pdf(docx_path, out_dir)
        if pdf:
            print(f"PDF:  {pdf}")
        else:
            print("PDF:  FAILED - LibreOffice (soffice) not available; deliver the DOCX instead.")


if __name__ == "__main__":
    main()
